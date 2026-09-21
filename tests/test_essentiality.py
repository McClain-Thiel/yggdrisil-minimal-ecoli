from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from yggdrisil_ecoli.data.errors import DataValidationError
from yggdrisil_ecoli.data.essentiality import (
    EssentialityRecord,
    parse_choe_workbook,
)
from yggdrisil_ecoli.data.evidence import load_genes, write_genes
from yggdrisil_ecoli.data.gff import parse_ncbi_gff
from yggdrisil_ecoli.scorers.essentiality import EssentialityScorer
from yggdrisil_ecoli.state import GenomeState

FIXTURES = Path(__file__).parent / "fixtures"


def test_parser_preserves_calls_conflicts_and_coordinate_audit(tmp_path: Path) -> None:
    registry, _ = parse_ncbi_gff(FIXTURES / "mg1655_excerpt.gff3")
    workbook_path = tmp_path / "choe.xlsx"
    _write_choe_fixture(workbook_path, mismatch_b0002=True)

    dataset, report = parse_choe_workbook(
        workbook_path, registry, expected_source_counts=None
    )

    assert dataset.loc["b0001"].classification == "essential"
    conditional = dataset.loc["b0002"]
    assert conditional.classification == "conditionally_essential"
    assert (conditional.lb_call_raw, conditional.m9_call_raw) == ("NE", "E")
    assert (conditional.lb_ecipkm, conditional.m9_ecipkm) == (10.0, 1.0)
    assert bool(conditional.condition_disagreement) is True
    assert bool(conditional.evidence_conflict) is False
    ambiguous = dataset.loc["b0003"]
    assert ambiguous.classification == "ambiguous"
    assert bool(ambiguous.evidence_conflict) is True
    assert report["unmapped_source_ids"] == ("b9999",)
    assert report["coordinate_mismatches"] == ("b0002",)


@pytest.mark.asyncio
async def test_one_table_round_trip_and_scorer_keep_unknown_separate(
    tmp_path: Path,
) -> None:
    registry, _ = parse_ncbi_gff(FIXTURES / "mg1655_excerpt.gff3")
    workbook_path = tmp_path / "choe.xlsx"
    _write_choe_fixture(workbook_path, omit_b0003=True)
    dataset, report = parse_choe_workbook(
        workbook_path,
        registry,
        expected_source_counts=None,
        metadata={"provenance": {"workbook_sha256": "fixture"}},
    )
    artifact = tmp_path / "genes.parquet"
    genes = registry.join(dataset, validate="one_to_one")
    genes.attrs.update(dataset.attrs)
    write_genes(genes, artifact)
    loaded = load_genes(artifact)

    assert len(loaded) == len(registry)
    assert report["canonical_genes_without_measurement"] == ("b0003",)
    assert loaded.loc["b0003", "classification"] == "unknown"
    assert loaded.loc["b0003", "coverage"] == "unknown"
    assert pd.isna(loaded.loc["b0003", "lb_ecipkm"])
    detail = loaded.loc["b0002"]
    assert detail["classification"] == "conditionally_essential"
    assert (detail["lb_call_raw"], detail["lb_ecipkm"]) == ("NE", 10.0)
    assert (detail["m9_call_raw"], detail["m9_ecipkm"]) == ("E", 1.0)
    assert loaded.attrs["essentiality"]["study_id"] == "choe2023_tnseq"
    assert loaded.attrs["essentiality"]["provenance"] == {"workbook_sha256": "fixture"}

    scorer = EssentialityScorer(genes=loaded, artifact_hash="artifact")
    result = await scorer.evaluate(GenomeState(frozenset({"b0001", "b0002", "b0003"})))

    assert result.metrics == {
        "n_essential_deleted": 1,
        "n_conditional_essential_deleted": 1,
        "n_ambiguous_deleted": 0,
        "n_unknown_deleted": 1,
    }
    assert result.metadata["details"] == {
        "essential_deleted": ["b0001"],
        "conditional_essential_deleted": ["b0002"],
        "ambiguous_deleted": [],
        "unknown_deleted": ["b0003"],
    }
    assert result.metadata["coverage"] == {
        "deleted_genes_total": 3,
        "deleted_genes_classified": 2,
        "deleted_genes_unknown": 1,
    }


def test_record_rejects_author_call_threshold_disagreement() -> None:
    with pytest.raises(ValidationError, match="disagrees with ecIPKM"):
        EssentialityRecord(
            b_number="b0001",
            lb_call_raw="E",
            lb_ecipkm=10.0,
            m9_call_raw="E",
            m9_ecipkm=1.0,
        )


@pytest.mark.parametrize(
    ("cell", "value", "reason"),
    [
        ("F5", "wrong", "b_number"),
        ("O5", "unexpected", "lb_call_raw"),
        ("N5", 1.0, "disagrees with ecIPKM"),
    ],
)
def test_parser_reports_worksheet_row_after_filtering(
    tmp_path: Path, cell: str, value: object, reason: str
) -> None:
    registry, _ = parse_ncbi_gff(FIXTURES / "mg1655_excerpt.gff3")
    path = tmp_path / "choe.xlsx"
    _write_choe_fixture(path)
    workbook = load_workbook(path)
    worksheet = workbook["Table S1"]
    worksheet.insert_rows(4)  # An excluded row before the now-fifth-row thrA record.
    worksheet["G4"] = "N"
    worksheet[cell] = value
    workbook.save(path)
    workbook.close()

    with pytest.raises(DataValidationError, match=f"(?s)row 5: .*{reason}"):
        parse_choe_workbook(path, registry, expected_source_counts=None)


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -1.0])
def test_nonfinite_or_negative_measurements_are_not_nonessential(value: float) -> None:
    with pytest.raises(ValidationError):
        EssentialityRecord(
            b_number="b0001",
            lb_call_raw="NE",
            lb_ecipkm=value,
            m9_call_raw="NE",
            m9_ecipkm=10.0,
        )


@pytest.mark.parametrize(
    "measurements",
    [
        {"lb_call_raw": "E", "lb_ecipkm": 1.0},
        {"lb_call_raw": "E", "lb_ecipkm": 1.0, "m9_call_raw": "NE"},
        {"lb_ecipkm": 1.0, "m9_ecipkm": 3.0},
    ],
)
def test_partial_measurements_are_rejected(measurements: dict[str, object]) -> None:
    with pytest.raises(ValidationError, match="both LB and M9"):
        EssentialityRecord.model_validate({"b_number": "b0001", **measurements})


def _write_choe_fixture(
    path: Path, *, omit_b0003: bool = False, mismatch_b0002: bool = False
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Table S1"
    worksheet.append(
        [
            "Gene",
            "Start",
            "End",
            "Length (nt)",
            "Strand",
            "Locus Tag",
            "CDS",
            "Pseudo",
            "PEC",
            "Gerdes",
            "LB medium",
            None,
            None,
            None,
            None,
            "M9 glucose (0.2%) medium",
            None,
            None,
            None,
            None,
        ]
    )
    worksheet.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Insertion",
            "IPKM",
            "ec Insertion",
            "ecIPKM",
            "Essentiality",
            "Insertion",
            "IPKM",
            "ec Insertion",
            "ecIPKM",
            "Essentiality",
        ]
    )
    worksheet.append(_source_row("thrL", "b0001", 190, 255, "E", "E"))
    worksheet.append(
        _source_row(
            "thrA",
            "b0002",
            338 if mismatch_b0002 else 337,
            2799,
            "NE",
            "E",
        )
    )
    if not omit_b0003:
        worksheet.append(_source_row("thrB", "b0003", 2801, 3733, "E", "NE"))
    worksheet.append(_source_row("outside", "b9999", 1, 3, "NE", "NE"))
    workbook.save(path)


def _source_row(
    symbol: str,
    b_number: str,
    start: int,
    end: int,
    lb_call: str,
    m9_call: str,
) -> list[object]:
    return [
        symbol,
        start,
        end,
        end - start + 1,
        "+",
        b_number,
        "Y",
        "N",
        "NE",
        "NE",
        1,
        1.0,
        1,
        1.0 if lb_call == "E" else 10.0,
        lb_call,
        1,
        1.0,
        1,
        1.0 if m9_call == "E" else 10.0,
        m9_call,
    ]
